"""Tests for the database manager (validation, import, rollback, access)."""
import io
import tempfile
from pathlib import Path

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import Workbook

from common.exceptions import ValidationError
from utils.excel import read_rows

from apps.database_manager.models import (
    DatabaseVersion,
    LabourMaster,
    RateMaster,
    StateControl,
    TORMain,
)
from apps.database_manager.services.importer import DatabaseImportService
from apps.database_manager.services.rollback import DatabaseRollbackService
from apps.database_manager.services.validator import validate_workbook

User = get_user_model()

REQUIRED_SHEETS = {
    "Rate_Master": (
        ["Tech_Key", "Match_Key", "Make", "Supplier", "Net Material Rate",
         "Unit", "Category", "Sub_Category"],
        [["P-100", "150 NB MS Pipe", "TATA", "VendorA", 1200, "m", "Pipes", "MS"]],
    ),
    "Labour_Master": (
        ["Tech_Key", "Labour_Type", "Total_Labour_with_Multiplier", "Unit"],
        [["L-1", "Fitter", 500, "day"]],
    ),
    "TOR_Main": (["Category"], [["Pipes"]]),
    "Labour_Structure_Source": (
        ["category", "sub_category", "size", "unit", "tech_key"],
        [["Pipes", "MS", "150", "m", "P-100"]],
    ),
    "TOR_Labour": (["Testing_%", "Scaffolding_%"], [[5, 10]]),
    "TOR_Accessories": (["Category", "Sub_Category", "Min_Size", "Max_Size", "Accessories_%"], [["Pipes", "MS", 0, 200, 5]]),
    "State_Control_List": (
        ["State", "Labour_Multiplier"],
        [["Maharashtra", "1.1"]],
    ),
}


def build_workbook(skip_sheet: str | None = None) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_name, (headers, rows) in REQUIRED_SHEETS.items():
        if sheet_name == skip_sheet:
            continue
        ws = wb.create_sheet(title=sheet_name)
        ws.append(headers)
        for row in rows:
            ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_client_database_workbook() -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Rate_Master")
    ws.append([
        "Column1", "Category", "Sub Category", "Class", "Size_mm", "Make",
        "Unit", "Supplier", "Base_Purchase_Rate", "Discount %",
        "Net Material Rate", "Handling_%", "Profit_%", "Final_Amount_(Excl GST)", "Tech_Key",
    ])
    ws.append([None, "PIPE", "MS", "C", 400, "TATA", "mm", "Tiger", 10977, 44, None, 0.02, 0.1, 7200, None])

    ws = wb.create_sheet("Labour_Master")
    ws.append(["Tech_Key", "Category", "Sub_Category", "Size", "Unit", "Base_Rate"])
    ws.append(["PIPE_MS_C_400", "PIPE", "MS", 400, "mm", 100])

    ws = wb.create_sheet("TOR_Main")
    ws.append(["Category  ", " Handling_%  ", "Wastage_%", "Profit_%  ", "Project_State"])
    ws.append(["PIPE", 0.02, 0.01, 0.1, "Delhi"])

    ws = wb.create_sheet("Labour_Structure_Source")
    ws.append(["Category", "Sub_Category", "Size", "Unit", "Tech_Key"])
    ws.append(["PIPE", "MS", 400, "mm", "PIPE_MS_C_400"])

    ws = wb.create_sheet("TOR_Labour")
    ws.append(["Testing_%", "Scaffolding_%", "Consumables_%", "Painting_Rate", "Labour_Buffer_%"])
    ws.append([0.05, 0.1, 0.05, 55, 0.04])

    ws = wb.create_sheet("TOR_Accessories")
    ws.append(["Category ", "Sub_Category ", "Min_Size", "Max_Size", "Accessories_% "])
    ws.append(["PIPE", "MS", 200, 400, 0.15])

    ws = wb.create_sheet("State_Control_List")
    ws.append(["State", "Labour_Multiplier"])
    ws.append(["Delhi", 1.2])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def write_temp_workbook(skip_sheet: str | None = None) -> str:
    data = build_workbook(skip_sheet)
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp.write(data)
    tmp.close()
    return tmp.name


class ValidatorTests(TestCase):
    def test_valid_workbook_passes(self):
        path = write_temp_workbook()
        try:
            sheets = validate_workbook(path)
            self.assertIn("Rate_Master", sheets)
        finally:
            Path(path).unlink(missing_ok=True)

    def test_missing_sheet_raises(self):
        path = write_temp_workbook(skip_sheet="Labour_Master")
        try:
            with self.assertRaises(ValidationError):
                validate_workbook(path)
        finally:
            Path(path).unlink(missing_ok=True)


class ReadRowsTests(TestCase):
    def test_read_rows_normalizes_headers(self):
        path = write_temp_workbook()
        try:
            rows = read_rows(path, "Rate_Master")
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0]["tech_key"], "P-100")
        finally:
            Path(path).unlink(missing_ok=True)


@override_settings(OPENAI_API_KEY="placeholder-key")
class ImportServiceTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser("admin@example.com", "pass12345")

    def _import(self) -> DatabaseVersion:
        path = write_temp_workbook()
        try:
            return DatabaseImportService(path, self.admin, "march.xlsx").run()
        finally:
            Path(path).unlink(missing_ok=True)

    def test_import_creates_active_version_and_rows(self):
        version = self._import()
        self.assertEqual(version.version_number, 1)
        self.assertTrue(version.is_active)
        self.assertEqual(RateMaster.objects.filter(database_version=version).count(), 1)
        self.assertEqual(LabourMaster.objects.filter(database_version=version).count(), 1)
        self.assertEqual(TORMain.objects.filter(database_version=version).count(), 1)
        self.assertEqual(StateControl.objects.filter(database_version=version).count(), 1)
        rate = RateMaster.objects.get(database_version=version)
        self.assertEqual(rate.tech_key, "P-100")
        self.assertEqual(str(rate.net_material_rate), "1200.00")

    def test_second_import_increments_and_deactivates_previous(self):
        first = self._import()
        second = self._import()
        first.refresh_from_db()
        self.assertEqual(second.version_number, 2)
        self.assertTrue(second.is_active)
        self.assertFalse(first.is_active)

    def test_retention_keeps_three_versions(self):
        """Import service retains active + two rollback versions."""
        for _ in range(5):
            self._import()
        self.assertEqual(DatabaseVersion.objects.count(), 3)
        numbers = sorted(DatabaseVersion.objects.values_list("version_number", flat=True))
        self.assertEqual(numbers, [3, 4, 5])

    def test_state_control_is_version_scoped(self):
        first = self._import()
        second = self._import()
        self.assertEqual(
            StateControl.objects.filter(
                database_version=first, state="Maharashtra"
            ).count(),
            1,
        )
        self.assertEqual(
            StateControl.objects.filter(
                database_version=second, state="Maharashtra"
            ).count(),
            1,
        )

    def test_imports_client_database_shape_with_synthesized_codes(self):
        tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        tmp.write(build_client_database_workbook())
        tmp.close()
        try:
            version = DatabaseImportService(tmp.name, self.admin, "client.xlsx").run()
        finally:
            Path(tmp.name).unlink(missing_ok=True)

        rate = RateMaster.objects.get(database_version=version)
        self.assertEqual(rate.tech_key, "PIPE_MS_C_400")
        self.assertEqual(rate.supplier, "Tiger")
        self.assertEqual(str(rate.net_material_rate), "6147.12")
        self.assertEqual(str(rate.final_amount_excl_gst), "7200.00")
        self.assertEqual(TORMain.objects.get(database_version=version).category, "PIPE")
        self.assertEqual(
            str(
                StateControl.objects.get(
                    database_version=version, state="Delhi"
                ).labour_multiplier
            ),
            "1.2000",
        )


@override_settings(OPENAI_API_KEY="placeholder-key")
class RollbackServiceTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser("admin@example.com", "pass12345")

    def _import(self) -> DatabaseVersion:
        path = write_temp_workbook()
        try:
            return DatabaseImportService(path, self.admin, "f.xlsx").run()
        finally:
            Path(path).unlink(missing_ok=True)

    def test_rollback_activates_target(self):
        first = self._import()
        second = self._import()
        first.refresh_from_db()  # mirror the view fetching a fresh instance
        DatabaseRollbackService(first).run()
        first.refresh_from_db()
        second.refresh_from_db()
        self.assertTrue(first.is_active)
        self.assertFalse(second.is_active)


@override_settings(MEDIA_ROOT=tempfile.mkdtemp(), OPENAI_API_KEY="placeholder-key")
class DatabaseAccessTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser("admin@example.com", "pass12345")
        self.expert = User.objects.create_user("expert@example.com", "pass12345")

    def test_expert_cannot_access_database(self):
        """Expert (without allow_db_access flag) must be denied access to upload.

        DatabaseVersionListView requires only LoginRequiredMixin (read-only listing
        is intentionally accessible to any logged-in user). The restricted actions
        (upload and rollback) are protected by DatabaseAccessRequiredMixin.
        """
        self.client.force_login(self.expert)
        self.assertEqual(
            self.client.get(reverse("database:upload")).status_code, 403
        )

    def test_super_admin_can_access_database(self):
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(reverse("database:list")).status_code, 200)

    def test_upload_imports_database(self):
        """Admin can upload a workbook and the import runs synchronously."""
        self.client.force_login(self.admin)
        upload = SimpleUploadedFile(
            "master.xlsx",
            build_workbook(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response = self.client.post(
            reverse("database:upload"),
            {"workbook": upload, "name": "Test Database v1"},
        )
        self.assertRedirects(response, reverse("database:list"))
        self.assertEqual(DatabaseVersion.objects.count(), 1)
        self.assertTrue(DatabaseVersion.objects.first().is_active)
