"""Tests for the export system (Sprint 19)."""
import io
import tempfile
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import Workbook
from openpyxl import load_workbook

from apps.boq.models import BOQ, BOQItem, BOQRun
from apps.costing.services.rate_detail import RateDetailRetrievalService
from apps.database_manager.models import DatabaseVersion, RateMaster
from apps.exports.models import ExportFile
from apps.exports.services.export_service import ExportService
from apps.matching.models import ProductMatch
from common.choices import BOQStatus, RunStatus
from common.exceptions import ValidationError

_MEDIA = tempfile.mkdtemp()


def _workbook_upload(name="source_boq.xlsx", rows=None):
    wb = Workbook()
    ws = wb.active
    rows = rows or [
        ["Particulars", "UOM", "Qty", "Remarks"],
        ["150 NB MS Pipe", "m", 3, "Existing"],
    ]
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return SimpleUploadedFile(
        name,
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@override_settings(MEDIA_ROOT=_MEDIA)
class ExportServiceTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(email="ex@x.com", password="x")
        self.version = DatabaseVersion.objects.create(
            version_number=1, is_active=True, source_filename="db.xlsx"
        )
        self.rate = RateMaster.objects.create(
            database_version=self.version,
            tech_key="PIPE150",
            make="APL",
            supplier="V1",
            net_material_rate=Decimal("800.00"),
            base_purchase_rate=Decimal("800.00"),
            final_amount_excl_gst=Decimal("800.00"),
        )
        self.boq = BOQ.objects.create(
            user=self.user, boq_name="B", uploaded_file="boq/x.xlsx", status=BOQStatus.APPROVED
        )
        self.run = BOQRun.objects.create(boq=self.boq, run_number=1, status=RunStatus.COMPLETED)
        self.item = BOQItem.objects.create(
            boq_run=self.run, row_number=1, description="150 NB MS Pipe", quantity=Decimal("3")
        )
        self.match = ProductMatch.objects.create(
            boq_item=self.item, product=self.rate, make="APL", supplier="V1",
            confidence_score=95, match_reason="exact", quantity_basis="per_boq_unit",
        )
        RateDetailRetrievalService().retrieve_item(self.match)

    def test_export_creates_files_and_marks_exported(self):
        export = ExportService().export_run(self.run, self.user)
        self.assertIsInstance(export, ExportFile)
        self.assertTrue(export.internal_sheet.name)
        self.assertTrue(export.client_sheet.name)
        self.boq.refresh_from_db()
        self.assertEqual(self.boq.status, BOQStatus.EXPORTED)

    def test_internal_workbook_has_both_sheets(self):
        export = ExportService().export_run(self.run, self.user)
        wb = load_workbook(io.BytesIO(export.internal_sheet.read()))
        self.assertEqual(wb.sheetnames, ["Breakdown List", "Client BOQ"])
        ws = wb["Breakdown List"]
        self.assertEqual(ws["A1"].value, "BOQ Ser No")
        self.assertEqual(ws["H2"].value, "PIPE150")  # product code row 1
        self.assertEqual(ws["Z1"].value, "Final_Amount_(Excl GST)")

    def test_internal_workbook_has_one_row_per_product_match(self):
        valve = RateMaster.objects.create(
            database_version=self.version,
            tech_key="VALVE80",
            make="Zoloto",
            supplier="V2",
            net_material_rate=Decimal("500.00"),
            base_purchase_rate=Decimal("500.00"),
            final_amount_excl_gst=Decimal("500.00"),
        )
        valve_match = ProductMatch.objects.create(
            boq_item=self.item,
            product=valve,
            make="Zoloto",
            supplier="V2",
            confidence_score=90,
            match_reason="exact",
            extraction_index=1,
            quantity_basis="per_boq_unit",
        )
        RateDetailRetrievalService().retrieve_item(valve_match)

        export = ExportService().export_run(self.run, self.user)
        wb = load_workbook(io.BytesIO(export.internal_sheet.read()))
        ws = wb["Breakdown List"]

        self.assertEqual(ws["H2"].value, "PIPE150")
        self.assertEqual(ws["H3"].value, "VALVE80")

    def test_client_standalone_has_computed_amount(self):
        export = ExportService().export_run(self.run, self.user)
        wb = load_workbook(io.BytesIO(export.client_sheet.read()))
        ws = wb["Client BOQ"]
        # Amount = final_rate * qty (computed, not a formula).
        final_rate = float(self.match.rate_detail.rate_contribution)
        self.assertAlmostEqual(ws["F2"].value, round(final_rate * 3, 2), places=2)

    def test_client_standalone_sums_multiple_product_rates(self):
        valve = RateMaster.objects.create(
            database_version=self.version,
            tech_key="VALVE80",
            make="Zoloto",
            supplier="V2",
            net_material_rate=Decimal("500.00"),
            base_purchase_rate=Decimal("500.00"),
            final_amount_excl_gst=Decimal("500.00"),
        )
        valve_match = ProductMatch.objects.create(
            boq_item=self.item,
            product=valve,
            confidence_score=90,
            match_reason="exact",
            extraction_index=1,
            quantity_basis="per_boq_unit",
        )
        RateDetailRetrievalService().retrieve_item(valve_match)

        export = ExportService().export_run(self.run, self.user)
        wb = load_workbook(io.BytesIO(export.client_sheet.read()))
        ws = wb["Client BOQ"]

        combined_rate = float(self.match.rate_detail.rate_contribution) + float(
            valve_match.rate_detail.rate_contribution
        )
        self.assertAlmostEqual(ws["E2"].value, combined_rate, places=2)
        self.assertAlmostEqual(ws["F2"].value, round(combined_rate * 3, 2), places=2)

    def test_client_standalone_uses_static_boq_columns(self):
        self.run.original_headers = [
            {"key": "ignored", "label": "Ignored", "index": 0},
            {"key": "description", "label": "Description", "index": 1},
            {"key": "quantity", "label": "Quantity", "index": 2},
            {"key": "unit", "label": "Unit", "index": 3},
        ]
        self.run.save(update_fields=["original_headers"])
        self.item.original_data = {
            "s_no": None,
            "description": "150 NB MS Pipe",
            "quantity": 3,
            "unit": "m",
        }
        self.item.save(update_fields=["original_data"])

        export = ExportService().export_run(self.run, self.user)
        wb = load_workbook(io.BytesIO(export.client_sheet.read()))
        ws = wb["Client BOQ"]

        self.assertEqual(ws["A1"].value, "S No")
        self.assertIsNone(ws["A2"].value)
        self.assertEqual(ws["B2"].value, "150 NB MS Pipe")
        self.assertEqual(ws["C1"].value, "Unit")
        self.assertEqual(ws["D1"].value, "Quantity")
        self.assertEqual(ws["E1"].value, "Final Rate")
        self.assertEqual(ws["F1"].value, "Amount")

    def test_client_export_preserves_uploaded_boq_and_fills_only_rate_amount_when_unit_qty_exist(self):
        self.boq.uploaded_file = _workbook_upload()
        self.boq.save(update_fields=["uploaded_file"])
        self.item.row_number = 2
        self.item.unit = "m"
        self.item.quantity = Decimal("3")
        self.item.row_json = {
            "rows": [
                {
                    "excel_row_number": 2,
                    "description": "150 NB MS Pipe",
                    "unit": "m",
                    "quantity": 3,
                }
            ]
        }
        self.item.save(update_fields=["row_number", "unit", "quantity", "row_json"])

        export = ExportService().export_run(self.run, self.user)
        wb = load_workbook(io.BytesIO(export.client_sheet.read()))
        ws = wb["Client BOQ"]

        self.assertEqual(ws["A1"].value, "Particulars")
        self.assertEqual(ws["B1"].value, "UOM")
        self.assertEqual(ws["C1"].value, "Qty")
        self.assertEqual(ws["D1"].value, "Remarks")
        self.assertEqual(ws["B2"].value, "m")
        self.assertEqual(ws["C2"].value, 3)
        self.assertEqual(ws["E1"].value, "Rate")
        self.assertEqual(ws["F1"].value, "Amount")
        self.assertEqual(ws["E2"].value, float(self.match.rate_detail.rate_contribution))
        self.assertEqual(ws["F2"].value, round(float(self.match.rate_detail.rate_contribution) * 3, 2))

    def test_client_export_fills_child_row_when_child_has_unit_quantity(self):
        self.boq.uploaded_file = _workbook_upload(
            rows=[
                ["Particulars", "UOM", "Qty"],
                ["Fire fighting pipe scope", None, None],
                ["150 NB MS Pipe", "m", 3],
            ]
        )
        self.boq.save(update_fields=["uploaded_file"])
        self.item.row_number = 2
        self.item.row_json = {
            "rows": [
                {"excel_row_number": 2, "description": "Fire fighting pipe scope"},
                {
                    "excel_row_number": 3,
                    "description": "150 NB MS Pipe",
                    "unit": "m",
                    "quantity": 3,
                },
            ]
        }
        self.item.save(update_fields=["row_number", "row_json"])

        export = ExportService().export_run(self.run, self.user)
        wb = load_workbook(io.BytesIO(export.client_sheet.read()))
        ws = wb["Client BOQ"]

        self.assertIsNone(ws["D2"].value)
        self.assertIsNone(ws["E2"].value)
        self.assertEqual(ws["D3"].value, float(self.match.rate_detail.rate_contribution))
        self.assertEqual(ws["E3"].value, round(float(self.match.rate_detail.rate_contribution) * 3, 2))

    def test_linked_client_sheet_references_internal_rate_and_client_quantity(self):
        export = ExportService().export_run(self.run, self.user)
        wb = load_workbook(io.BytesIO(export.internal_sheet.read()), data_only=False)
        ws = wb["Client BOQ"]

        self.assertEqual(ws["E2"].value, "='Breakdown List'!AA2")
        self.assertEqual(ws["F2"].value, "='Breakdown List'!AA2*D2")

    def test_linked_client_sheet_sums_multiple_internal_rows(self):
        valve = RateMaster.objects.create(
            database_version=self.version,
            tech_key="VALVE80",
            net_material_rate=Decimal("500.00"),
            base_purchase_rate=Decimal("500.00"),
            final_amount_excl_gst=Decimal("500.00"),
        )
        valve_match = ProductMatch.objects.create(
            boq_item=self.item,
            product=valve,
            confidence_score=90,
            match_reason="exact",
            extraction_index=1,
            quantity_basis="per_boq_unit",
        )
        RateDetailRetrievalService().retrieve_item(valve_match)

        export = ExportService().export_run(self.run, self.user)
        wb = load_workbook(io.BytesIO(export.internal_sheet.read()), data_only=False)
        ws = wb["Client BOQ"]

        self.assertEqual(ws["E2"].value, "=SUM('Breakdown List'!AA2,'Breakdown List'!AA3)")
        self.assertEqual(ws["F2"].value, "=SUM('Breakdown List'!AA2,'Breakdown List'!AA3)*D2")

    def test_export_requires_approved(self):
        self.boq.status = BOQStatus.COMPLETED
        self.boq.save(update_fields=["status"])
        with self.assertRaises(ValidationError):
            ExportService().export_run(self.run, self.user)


@override_settings(MEDIA_ROOT=_MEDIA)
class ExportViewTests(TestCase):
    def setUp(self):
        self.owner = get_user_model().objects.create_user(email="o@x.com", password="x")
        self.other = get_user_model().objects.create_user(email="t@x.com", password="x")
        self.version = DatabaseVersion.objects.create(
            version_number=1, is_active=True, source_filename="db.xlsx"
        )
        self.rate = RateMaster.objects.create(
            database_version=self.version,
            tech_key="PIPE150",
            net_material_rate=Decimal("800.00"),
            base_purchase_rate=Decimal("800.00"),
            final_amount_excl_gst=Decimal("800.00"),
        )
        self.boq = BOQ.objects.create(
            user=self.owner, boq_name="B", uploaded_file="boq/x.xlsx", status=BOQStatus.APPROVED
        )
        self.run = BOQRun.objects.create(boq=self.boq, run_number=1, status=RunStatus.COMPLETED)
        item = BOQItem.objects.create(
            boq_run=self.run, row_number=1, description="150 NB MS Pipe", quantity=Decimal("1")
        )
        match = ProductMatch.objects.create(
            boq_item=item, product=self.rate, confidence_score=95, match_reason="exact",
            quantity_basis="per_boq_unit",
        )
        RateDetailRetrievalService().retrieve_item(match)

    def test_owner_can_export(self):
        self.client.force_login(self.owner)
        resp = self.client.post(reverse("exports:generate", args=[self.boq.pk]))
        self.assertEqual(resp.status_code, 302)
        self.boq.refresh_from_db()
        self.assertEqual(self.boq.status, BOQStatus.EXPORTED)
        self.assertTrue(ExportFile.objects.filter(boq_run=self.run).exists())

    def test_owner_can_download_exported_files(self):
        ExportService().export_run(self.run, self.owner)
        self.client.force_login(self.owner)

        client_resp = self.client.get(reverse("exports:download", args=[self.boq.pk, "client"]))
        breakdown_resp = self.client.get(
            reverse("exports:download", args=[self.boq.pk, "breakdown"])
        )

        self.assertEqual(client_resp.status_code, 200)
        self.assertEqual(breakdown_resp.status_code, 200)
        self.assertIn("attachment", client_resp["Content-Disposition"])
        self.assertIn("attachment", breakdown_resp["Content-Disposition"])

    def test_non_owner_cannot_export(self):
        self.client.force_login(self.other)
        resp = self.client.post(reverse("exports:generate", args=[self.boq.pk]))
        self.assertEqual(resp.status_code, 404)
        self.assertFalse(ExportFile.objects.exists())
