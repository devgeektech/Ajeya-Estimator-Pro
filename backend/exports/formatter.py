"""Shared Excel formatting helpers for exports.

Header styling + confidence colour bands (docs/PRD.md - Confidence Scoring).
"""
from __future__ import annotations

from openpyxl.styles import Alignment, Font, PatternFill

CONFIDENCE_COLORS = {
    "green": "C6EFCE",
    "yellow": "FFEB9C",
    "orange": "FCE4D6",
    "red": "FFC7CE",
    "blank": "FFFFFF",
}

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="2563EB")
_CENTER = Alignment(horizontal="center", vertical="center")


def write_header(ws, headers: list[str]) -> None:
    """Write a styled header row and set basic column widths."""
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER
        ws.column_dimensions[cell.column_letter].width = max(12, len(title) + 2)
    ws.freeze_panes = "A2"


def band_fill(band: str) -> PatternFill:
    """PatternFill for a confidence band label."""
    return PatternFill("solid", fgColor=CONFIDENCE_COLORS.get(band, "FFFFFF"))
