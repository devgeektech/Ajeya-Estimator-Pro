"""Client sheet generation (Phase 10, Sprint 19).

Original BOQ format with final approved rates and amounts
(docs/PRD.md - Client Sheet). When generated alongside the internal sheet the
rate/amount cells are Excel formulas linked to the internal sheet so edits there
flow through ("Values are linked to the internal sheet"); as a standalone client
file the values are computed.
"""
from __future__ import annotations

from copy import copy

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from apps.boq.services.parser import CANONICAL_BOQ_HEADERS
from apps.boq.services.parser import QUANTITY_KEYS, UNIT_KEYS
from utils.excel import _normalize_header

from .formatter import write_header
from .internal_sheet import FINAL_RATE_COL

RATE_KEYS = (
    "rate", "final_rate", "unit_rate", "quoted_rate", "approved_rate",
    "basic_rate", "rate_in_rs", "rate_rs",
)
AMOUNT_KEYS = (
    "amount", "final_amount", "total_amount", "total", "amt",
    "final_amount_excl_gst",
)

def _num(value) -> float:
    if value in (None, ""):
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _base_headers(run) -> list[dict]:
    return CANONICAL_BOQ_HEADERS


def _cell_value(item, header: dict):
    key = header["key"]
    if key == "row_number":
        return item.row_number
    if key == "s_no":
        return (item.original_data or {}).get("s_no", "")
    if key == "description":
        return item.description
    if key == "quantity":
        return _num(item.quantity)
    if key == "unit":
        return item.unit
    value = (item.original_data or {}).get(key)
    return "" if value is None else value


def _copy_source_sheet(target_ws, run) -> bool:
    source_file = run.boq.uploaded_file
    if not source_file:
        return False
    try:
        source_wb = load_workbook(source_file.path, data_only=False)
    except Exception:
        return False
    try:
        source_ws = source_wb.active
        target_ws.title = "Client BOQ"
        for row in source_ws.iter_rows():
            for cell in row:
                target = target_ws.cell(row=cell.row, column=cell.column, value=cell.value)
                if cell.has_style:
                    target.font = copy(cell.font)
                    target.fill = copy(cell.fill)
                    target.border = copy(cell.border)
                    target.alignment = copy(cell.alignment)
                    target.number_format = cell.number_format
                    target.protection = copy(cell.protection)
                if cell.hyperlink:
                    target.hyperlink = copy(cell.hyperlink)
                if cell.comment:
                    target.comment = copy(cell.comment)
        for merged_range in source_ws.merged_cells.ranges:
            target_ws.merge_cells(str(merged_range))
        for key, dimension in source_ws.column_dimensions.items():
            target_ws.column_dimensions[key].width = dimension.width
            target_ws.column_dimensions[key].hidden = dimension.hidden
        for key, dimension in source_ws.row_dimensions.items():
            target_ws.row_dimensions[key].height = dimension.height
            target_ws.row_dimensions[key].hidden = dimension.hidden
        target_ws.freeze_panes = source_ws.freeze_panes
        return True
    finally:
        source_wb.close()


def _header_row_and_columns(ws) -> tuple[int | None, dict[str, int]]:
    candidates = []
    wanted = set(UNIT_KEYS + QUANTITY_KEYS + RATE_KEYS + AMOUNT_KEYS)
    for row_number in range(1, min(ws.max_row, 30) + 1):
        columns: dict[str, int] = {}
        matches = 0
        for column in range(1, ws.max_column + 1):
            key = _normalize_header(ws.cell(row_number, column).value)
            if not key:
                continue
            columns[key] = column
            if key in wanted:
                matches += 1
        if matches:
            candidates.append((matches, row_number, columns))
    if not candidates:
        return None, {}
    _, row_number, columns = max(candidates, key=lambda item: item[0])
    return row_number, columns


def _first_matching_column(columns: dict[str, int], keys) -> int | None:
    for key in keys:
        if key in columns:
            return columns[key]
    return None


def _ensure_column(ws, header_row: int, columns: dict[str, int], keys, label: str) -> int:
    existing = _first_matching_column(columns, keys)
    if existing:
        return existing
    column = ws.max_column + 1
    ws.cell(row=header_row, column=column, value=label)
    columns[_normalize_header(label)] = column
    ws.column_dimensions[get_column_letter(column)].width = max(12, len(label) + 2)
    return column


def _target_excel_row(item, unit_col: int | None, qty_col: int | None, ws) -> int:
    if item.target_excel_row:
        return item.target_excel_row
    rows = (item.row_json or {}).get("rows") or []
    for row in rows:
        excel_row = row.get("excel_row_number")
        if not excel_row:
            continue
        has_unit = unit_col and ws.cell(excel_row, unit_col).value not in (None, "")
        has_qty = qty_col and ws.cell(excel_row, qty_col).value not in (None, "")
        if has_unit or has_qty:
            return excel_row
    for row in rows:
        excel_row = row.get("excel_row_number")
        if excel_row:
            return excel_row
    return item.row_number


def _breakdown_rate(item) -> float:
    if any(match.review_required or match.product_id is None for match in item.product_matches.all()):
        return 0.0
    total = 0.0
    for match in item.product_matches.all():
        detail = getattr(match, "rate_detail", None)
        total += _num(detail.rate_contribution) if detail else 0.0
    return total


def _breakdown_rows_by_item(run) -> dict[int, list[int]]:
    """Return internal breakdown row numbers keyed by BOQItem id."""
    mapping: dict[int, list[int]] = {}
    row = 2
    for item in run.items.all().order_by("row_number"):
        matches = list(item.product_matches.all())
        row_count = max(1, len(matches))
        mapping[item.pk] = list(range(row, row + row_count))
        row += row_count
    return mapping


def _linked_rate_formula(link_sheet: str, final_rate_letter: str, rows: list[int]) -> str:
    refs = [f"'{link_sheet}'!{final_rate_letter}{row}" for row in rows]
    if not refs:
        return "0"
    if len(refs) == 1:
        return refs[0]
    return f"SUM({','.join(refs)})"


def _write_original_client_sheet(ws, run, *, link_sheet: str | None) -> bool:
    if not _copy_source_sheet(ws, run):
        return False
    header_row, columns = _header_row_and_columns(ws)
    if header_row is None:
        return False

    unit_col = _ensure_column(ws, header_row, columns, UNIT_KEYS, "Unit")
    qty_col = _ensure_column(ws, header_row, columns, QUANTITY_KEYS, "Quantity")
    rate_col = _ensure_column(ws, header_row, columns, RATE_KEYS, "Rate")
    amount_col = _ensure_column(ws, header_row, columns, AMOUNT_KEYS, "Amount")
    final_rate_letter = get_column_letter(FINAL_RATE_COL)
    breakdown_rows = _breakdown_rows_by_item(run)

    for item in run.items.all().order_by("row_number"):
        target_row = _target_excel_row(item, unit_col, qty_col, ws)
        if unit_col and ws.cell(target_row, unit_col).value in (None, ""):
            ws.cell(target_row, unit_col, value=item.unit)
        if qty_col and ws.cell(target_row, qty_col).value in (None, ""):
            ws.cell(target_row, qty_col, value=_num(item.quantity))

        qty_ref = f"{get_column_letter(qty_col)}{target_row}"
        if link_sheet:
            rate_ref = _linked_rate_formula(
                link_sheet,
                final_rate_letter,
                breakdown_rows.get(item.pk, []),
            )
            if _breakdown_rate(item):
                ws.cell(target_row, rate_col, value=f"={rate_ref}")
                ws.cell(target_row, amount_col, value=f"={rate_ref}*{qty_ref}")
            else:
                ws.cell(target_row, rate_col, value=None)
                ws.cell(target_row, amount_col, value=None)
        else:
            final_rate = _breakdown_rate(item)
            qty = _num(ws.cell(target_row, qty_col).value) or _num(item.quantity)
            ws.cell(target_row, rate_col, value=final_rate or None)
            ws.cell(target_row, amount_col, value=round(final_rate * qty, 2) if final_rate else None)
    return True


def write_client_sheet(ws, run, *, link_sheet: str | None = None) -> None:
    """Populate ``ws`` with the client BOQ.

    If ``link_sheet`` is the internal sheet's title, the Final Rate / Amount
    cells reference it via formulas; otherwise computed values are written.
    """
    if _write_original_client_sheet(ws, run, link_sheet=link_sheet):
        return None

    ws.title = "Client BOQ"
    base_headers = _base_headers(run)
    headers = [header["label"] for header in base_headers] + ["Final Rate", "Amount"]
    write_header(ws, headers)

    col_letter = get_column_letter(FINAL_RATE_COL)  # internal Final Rate column
    quantity_col = next(
        index + 1
        for index, header in enumerate(base_headers)
        if header["key"] == "quantity"
    )
    quantity_letter = get_column_letter(quantity_col)
    breakdown_rows = _breakdown_rows_by_item(run)
    row = 2
    for item in run.items.all().order_by("row_number"):
        qty = _num(item.quantity)
        final_rate = _breakdown_rate(item)

        for column, header in enumerate(base_headers, start=1):
            ws.cell(row=row, column=column, value=_cell_value(item, header))

        rate_col = len(base_headers) + 1
        amount_col = len(base_headers) + 2
        if link_sheet:
            ref = _linked_rate_formula(link_sheet, col_letter, breakdown_rows.get(item.pk, []))
            if final_rate:
                ws.cell(row=row, column=rate_col, value=f"={ref}")
                ws.cell(row=row, column=amount_col, value=f"={ref}*{quantity_letter}{row}")
            else:
                ws.cell(row=row, column=rate_col, value=None)
                ws.cell(row=row, column=amount_col, value=None)
        else:
            ws.cell(row=row, column=rate_col, value=final_rate or None)
            ws.cell(row=row, column=amount_col, value=round(final_rate * qty, 2) if final_rate else None)
        row += 1

    return None
