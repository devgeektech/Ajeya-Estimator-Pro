"""Convert legacy Excel 97-2003 (``.xls``) workbooks to Open XML (``.xlsx``).

Uses xlrd to read cell values and openpyxl to write a new workbook so the rest
of the pipeline can keep using openpyxl. Formatting, charts, and macros are
not preserved — only sheet names and cell values (including Excel dates).
"""
from __future__ import annotations

import logging
from contextlib import contextmanager
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Iterator

import xlrd
from openpyxl import Workbook
from xlrd import XL_CELL_BOOLEAN, XL_CELL_DATE, XL_CELL_ERROR, XL_CELL_NUMBER

logger = logging.getLogger("boq_ai")


class XlsConvertError(ValueError):
    """Raised when a ``.xls`` workbook cannot be converted."""


def is_xls_filename(name: str | None) -> bool:
    """True when ``name`` is a legacy ``.xls`` file (not ``.xlsx`` / ``.xlsm``)."""
    lower = (name or "").lower().rsplit("/", 1)[-1]
    return lower.endswith(".xls") and not lower.endswith((".xlsx", ".xlsm"))


def xls_to_xlsx_filename(name: str) -> str:
    """Replace a ``.xls`` suffix with ``.xlsx`` (preserves basename)."""
    basename = (name or "workbook.xls").rsplit("/", 1)[-1]
    if is_xls_filename(basename):
        return basename[:-4] + ".xlsx"
    if not basename.lower().endswith((".xlsx", ".xlsm")):
        return f"{basename}.xlsx"
    return basename


def _cell_value(book: xlrd.Book, cell):
    """Map an xlrd cell to a Python value openpyxl can store."""
    if cell.ctype == XL_CELL_DATE:
        try:
            return xlrd.xldate_as_datetime(cell.value, book.datemode)
        except Exception:
            return cell.value
    if cell.ctype == XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell.ctype == XL_CELL_ERROR:
        return None
    if cell.ctype == XL_CELL_NUMBER:
        # Prefer ints when the float is whole (common for S.No / qty).
        value = cell.value
        if isinstance(value, float) and value.is_integer():
            return int(value)
        return value
    return cell.value


def convert_xls_path_to_xlsx(source_path: str | Path, dest_path: str | Path) -> Path:
    """Convert ``source_path`` (``.xls``) into ``dest_path`` (``.xlsx``).

    Only **visible** worksheets are copied (hidden / very-hidden sheets are skipped).
    """
    source = Path(source_path)
    dest = Path(dest_path)
    try:
        book = xlrd.open_workbook(str(source), formatting_info=False)
    except Exception as exc:
        raise XlsConvertError(
            f"Could not read legacy Excel file '{source.name}'. "
            "Open it in Excel and Save As .xlsx, then upload again."
        ) from exc

    workbook = Workbook()
    default_sheet = workbook.active
    if default_sheet is not None:
        workbook.remove(default_sheet)

    try:
        for sheet_index in range(book.nsheets):
            sheet = book.sheet_by_index(sheet_index)
            # 0 = visible, 1 = hidden, 2 = very hidden
            if int(getattr(sheet, "visibility", 0) or 0) != 0:
                continue
            title = (sheet.name or f"Sheet{sheet_index + 1}")[:31]
            worksheet = workbook.create_sheet(title=title)
            for row_index in range(sheet.nrows):
                for col_index in range(sheet.ncols):
                    cell = sheet.cell(row_index, col_index)
                    if cell.ctype == xlrd.XL_CELL_EMPTY:
                        continue
                    worksheet.cell(
                        row=row_index + 1,
                        column=col_index + 1,
                        value=_cell_value(book, cell),
                    )
        if not workbook.sheetnames:
            workbook.create_sheet(title="Sheet1")
        dest.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(str(dest))
    except XlsConvertError:
        raise
    except Exception as exc:
        raise XlsConvertError(
            f"Could not convert '{source.name}' to .xlsx. "
            "Open it in Excel and Save As .xlsx, then upload again."
        ) from exc

    logger.info("Converted legacy Excel '%s' → '%s'", source.name, dest.name)
    return dest


def list_xls_visible_sheet_names(file_path: str | Path) -> list[str]:
    """Return visible worksheet names from a legacy ``.xls`` workbook."""
    try:
        book = xlrd.open_workbook(str(file_path), formatting_info=False)
    except Exception as exc:
        raise XlsConvertError(
            f"Could not read legacy Excel file '{Path(file_path).name}'."
        ) from exc
    names: list[str] = []
    for sheet_index in range(book.nsheets):
        sheet = book.sheet_by_index(sheet_index)
        if int(getattr(sheet, "visibility", 0) or 0) != 0:
            continue
        title = str(sheet.name or "").strip()
        if title:
            names.append(title)
    return names


@contextmanager
def openxml_workbook_path(file_path: str | Path) -> Iterator[str]:
    """Yield a path openpyxl can read; convert ``.xls`` to a temp ``.xlsx`` first."""
    path = Path(file_path)
    if not is_xls_filename(path.name):
        yield str(path)
        return

    tmp_path: str | None = None
    try:
        with NamedTemporaryFile(suffix=".xlsx", delete=False) as handle:
            tmp_path = handle.name
        convert_xls_path_to_xlsx(path, tmp_path)
        yield tmp_path
    finally:
        if tmp_path:
            Path(tmp_path).unlink(missing_ok=True)
