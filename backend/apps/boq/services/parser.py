"""BOQ and make-list parsing.

BOQ workbooks arrive in varying client formats, so header matching is tolerant
and falls back gracefully (docs/PRD.md - Product Matching Rules: understand
varying descriptions). Parsing only captures the original rows; AI
understanding and matching happen in later phases.
"""
from __future__ import annotations

import logging
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import re

from openpyxl import load_workbook

from utils.excel import _normalize_header
from utils.excel import read_rows_with_metadata

logger = logging.getLogger("boq_ai")

# Accepted header aliases (already normalized to lower snake_case).
CANONICAL_BOQ_HEADERS = [
    {"key": "s_no", "label": "S No"},
    {"key": "description", "label": "Description"},
    {"key": "unit", "label": "Unit"},
    {"key": "quantity", "label": "Quantity"},
]
DESCRIPTION_KEYS = (
    "description", "item_description", "item", "particulars",
    "work_description", "scope", "description_of_work", "description_of_item",
    "item_of_work", "brief_description", "boq_description", "details",
    "specification", "specifications", "desc", "item_desc", "particular",
)
QUANTITY_KEYS = (
    "quantity", "qty", "qnty", "quantities", "nos", "no", "total_qty",
    "estimated_qty", "qty_required", "required_qty",
)
UNIT_KEYS = (
    "unit", "uom", "units", "unit_of_measure", "measure", "un", "ut",
    "unit_measure",
)
SERIAL_KEYS = (
    "s_no", "sno", "sr_no", "srno", "serial_no", "sl_no", "slno", "item_no",
    "line_no",
)

MAKE_KEYS = (
    "make", "makes", "make_list", "approved_make", "approved_makes",
    "approved_brand", "approved_brands", "brand", "brands", "manufacturer",
    "manufacturers", "make_manufacturer", "make_manufacturers",
    "make_manufacturer_name", "make_manufacturers_name",
    "manufacturers_name", "manufacturer_name", "oem", "vendor_make",
    "make_1", "make_2", "make_3", "brand_1", "brand_2", "brand_3",
)
CATEGORY_KEYS = (
    "category", "item_category", "type", "section", "item", "item_desc",
    "item_description", "description", "material", "materials", "product",
    "equipment",
)
BOQ_HEADER_KEYS = set(DESCRIPTION_KEYS + QUANTITY_KEYS + UNIT_KEYS + SERIAL_KEYS)
MAKE_LIST_HEADER_KEYS = set(MAKE_KEYS + CATEGORY_KEYS)


def _pick(row: dict, keys) -> object:
    for key in keys:
        if key in row and row[key] not in (None, ""):
            return row[key]
    return None


def _to_decimal(value) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError):
        return Decimal("0")


def _json_safe(value):
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, datetime | date):
        return value.isoformat()
    return value


def _fallback_description(row: dict) -> str:
    for value in row.values():
        if value not in (None, ""):
            return str(value).strip()
    return ""


def _canonical_original_data(row: dict) -> dict:
    serial_no = _pick(row, SERIAL_KEYS)
    description = _pick(row, DESCRIPTION_KEYS)
    unit = _pick(row, UNIT_KEYS)
    quantity = _pick(row, QUANTITY_KEYS)
    return {
        "s_no": _json_safe(serial_no),
        "description": _json_safe(description),
        "unit": _json_safe(unit),
        "quantity": _json_safe(quantity),
    }


def _serial_value(item: dict):
    return (item.get("original_data") or {}).get("s_no")


def _has_serial(item: dict) -> bool:
    value = _serial_value(item)
    return value not in (None, "")


def _row_payload(group: list[dict], primary: dict | None = None) -> dict:
    primary = primary or group[0]
    rows = []
    descriptions = []
    for item in group:
        original = item.get("original_data") or {}
        description = str(original.get("description") or item.get("description") or "").strip()
        if description:
            descriptions.append(description)
        rows.append(
            {
                "excel_row_number": item["row_number"],
                "serial_number": original.get("s_no"),
                "description": description,
                "unit": original.get("unit") or item.get("unit", ""),
                "quantity": original.get("quantity"),
                "canonical": original,
            }
        )

    return {
        "schema": "boq_row_group_v1",
        "primary_excel_row_number": primary["row_number"],
        "target_excel_row": primary["row_number"],
        "excel_row_numbers": [item["row_number"] for item in group],
        "serial_number": _serial_value(primary),
        "description": "\n".join(descriptions),
        "unit": (primary.get("original_data") or {}).get("unit") or primary.get("unit", ""),
        "quantity": (primary.get("original_data") or {}).get("quantity"),
        "rows": rows,
    }


def _group_boq_items(items: list[dict]) -> list[dict]:
    """Attach parent/specification rows to the measured BOQ rows.

    If a workbook has no serial numbers at all, each parsed row remains its own
    BOQ item. When serial numbers are present, rows without unit/quantity are
    treated as context for the next measured row instead of being processed as
    separate priced items.
    """
    if not any(_has_serial(item) for item in items):
        groups = [[item] for item in items]
        return [_build_grouped_item(group, group[0]) for group in groups]

    measured_groups = _measured_boq_groups(items)
    if measured_groups:
        return [_build_grouped_item(group, primary) for group, primary in measured_groups]

    groups = [[item] for item in items]
    return [_build_grouped_item(group, group[0]) for group in groups]


def _build_grouped_item(group: list[dict], primary: dict) -> dict:
    payload = _row_payload(group, primary)
    return {
        "row_number": primary["row_number"],
        "target_excel_row": primary["row_number"],
        "description": payload["description"],
        "quantity": primary["quantity"],
        "unit": primary["unit"],
        "original_data": primary["original_data"],
        "row_json": payload,
    }


def _measured_boq_groups(items: list[dict]) -> list[tuple[list[dict], dict]]:
    groups: list[tuple[list[dict], dict]] = []
    context_by_level: dict[int, dict] = {}
    context_details: list[dict] = []
    variant_details: list[dict] = []
    previous_group: list[dict] | None = None

    for item in items:
        if _has_measurement(item):
            if _has_serial(item) and _serial_level(_serial_value(item)) <= 1:
                level = _serial_level(_serial_value(item))
                same_or_deeper_context = any(
                    existing_level >= level for existing_level in context_by_level
                )
                context_by_level = {
                    existing_level: context_item
                    for existing_level, context_item in context_by_level.items()
                    if existing_level < level
                }
                if same_or_deeper_context:
                    context_details = []
                    variant_details = []
            group = _current_context(context_by_level) + context_details + variant_details + [item]
            groups.append((group, item))
            previous_group = group
            variant_details = []
            continue

        if _has_serial(item):
            level = _serial_level(_serial_value(item))
            if level <= 1:
                context_by_level = {
                    existing_level: context_item
                    for existing_level, context_item in context_by_level.items()
                    if existing_level < level
                }
                context_by_level[level] = item
                context_details = []
                variant_details = []
            else:
                variant_details.append(item)
            continue

        if variant_details:
            variant_details.append(item)
        elif context_by_level:
            context_details.append(item)
        elif previous_group is not None:
            previous_group.append(item)
        else:
            context_details.append(item)

    return groups


def _current_context(context_by_level: dict[int, dict]) -> list[dict]:
    return [
        context_by_level[level]
        for level in sorted(context_by_level)
    ]


def _has_measurement(item: dict) -> bool:
    original = item.get("original_data") or {}
    return bool(original.get("unit")) or original.get("quantity") not in (None, "")


def _serial_level(value) -> int:
    text = str(value or "").strip()
    if not text:
        return 2
    normalized = text.strip("()").rstrip(")").strip()
    parenthesized = text.startswith("(") or text.endswith(")")
    if parenthesized and re.fullmatch(r"[A-Za-z]+", normalized):
        return 2
    if re.fullmatch(r"[IVXLCDM]+", normalized, flags=re.IGNORECASE):
        return 0
    if re.fullmatch(r"\d+(?:\.0+)?", normalized):
        return 0
    if re.fullmatch(r"\d+(?:\.\d+)+", normalized):
        return min(normalized.count("."), 1)
    if re.fullmatch(r"[A-Za-z]+", normalized):
        return 2
    return 1


def _split_make_names(value) -> list[str]:
    if value in (None, ""):
        return []
    text = str(value).strip()
    if not text:
        return []
    text = re.sub(r"\bM\s*/\s*s\b", "M_SLASH_S", text, flags=re.IGNORECASE)
    parts: list[str] = []
    for group in re.split(r"\s*(?:;|\n|\r|\|)\s*", text):
        if not group.strip():
            continue
        separator = r"\s*/\s*" if "/" in group else r"\s*,\s*"
        parts.extend(re.split(separator, group))
    return [
        part.replace("M_SLASH_S", "M/s").strip()
        for part in parts
        if part.replace("M_SLASH_S", "M/s").strip()
    ]


def _make_entry_key(make: str, category: str) -> tuple[str, str]:
    return (make.strip().casefold(), category.strip().casefold())


def _make_values_from_row(row: dict) -> list[str]:
    values: list[str] = []
    for key, value in row.items():
        if key in CATEGORY_KEYS:
            continue
        if key in MAKE_KEYS or key.startswith(
            ("make_", "brand_", "manufacturer_", "approved_make_")
        ):
            values.extend(_split_make_names(value))
    return values


def parse_boq_workbook(file_path: str) -> tuple[list[dict], list[dict]]:
    """Return canonical headers and captured BOQ rows from a workbook."""
    _, rows = read_rows_with_metadata(file_path, header_keys=BOQ_HEADER_KEYS)
    captured: list[dict] = []
    for record in rows:
        raw_row = record["values"]
        display_row = record["display_values"]
        row = {key: _json_safe(value) for key, value in display_row.items()}
        description = _pick(row, DESCRIPTION_KEYS)
        has_description_column = any(key in row for key in DESCRIPTION_KEYS)
        if description in (None, "") and not has_description_column:
            description = _fallback_description(row)
        if not description:
            continue
        captured.append(
            {
                "row_number": record["excel_row_number"],
                "description": str(description).strip(),
                "quantity": _to_decimal(_pick(raw_row, QUANTITY_KEYS)),
                "unit": (str(_pick(row, UNIT_KEYS)).strip() if _pick(row, UNIT_KEYS) else ""),
                "original_data": _canonical_original_data(row),
            }
        )
    return CANONICAL_BOQ_HEADERS, _group_boq_items(captured)


def parse_boq_items(file_path: str) -> list[dict]:
    """Return BOQ item dicts while preserving worksheet row numbers."""
    _, items = parse_boq_workbook(file_path)
    return items


def parse_make_list(file_path: str) -> list[dict]:
    """Return a list of {make, category} dicts from a make-list workbook or PDF."""
    if str(file_path).lower().endswith(".pdf"):
        return _parse_make_list_pdf(file_path)

    return _parse_make_list_workbook(file_path)


def _parse_make_list_workbook(file_path: str) -> list[dict]:
    """Extract make-list entries from all sheets in a workbook.

    Real client files may keep the actual make list in a later worksheet and
    may use a merged "Approved Makes" header spanning multiple make columns.
    """
    entries: list[dict] = []
    seen: set[tuple[str, str]] = set()
    workbook = load_workbook(filename=file_path, read_only=True, data_only=True)
    try:
        for worksheet in workbook.worksheets:
            _append_make_list_sheet_entries(worksheet, entries, seen)
    finally:
        workbook.close()
    return entries


def _append_make_list_sheet_entries(
    worksheet, entries: list[dict], seen: set[tuple[str, str]]
) -> None:
    rows = list(worksheet.iter_rows(values_only=True))
    header = _make_list_header(rows)
    if not header:
        return
    header_row_idx, category_cols, make_cols = header
    for row in rows[header_row_idx + 1:]:
        category = _first_row_value(row, category_cols)
        for make_col in make_cols:
            if make_col >= len(row):
                continue
            for make in _split_make_names(row[make_col]):
                key = _make_entry_key(make, category)
                if key in seen:
                    continue
                seen.add(key)
                entries.append({"make": make, "category": category})


def _make_list_header(rows: list[tuple]) -> tuple[int, list[int], list[int]] | None:
    best: tuple[int, list[int], list[int], int] | None = None
    for idx, row in enumerate(rows[:30]):
        normalized = [_normalize_header(value) for value in row]
        category_cols = [col for col, key in enumerate(normalized) if key in CATEGORY_KEYS]
        direct_make_cols = [col for col, key in enumerate(normalized) if _is_make_header(key)]
        if not direct_make_cols:
            continue
        make_cols = _expanded_make_columns(normalized, direct_make_cols)
        score = len(category_cols) + len(make_cols) * 2
        if best is None or score > best[3]:
            best = (idx, category_cols, make_cols, score)
    if best is None:
        return None
    return best[0], best[1], best[2]


def _is_make_header(key: str) -> bool:
    return key in MAKE_KEYS or key.startswith(
        ("make_", "brand_", "manufacturer_", "approved_make_")
    )


def _expanded_make_columns(normalized_headers: list[str], direct_make_cols: list[int]) -> list[int]:
    make_cols: list[int] = []
    for col in direct_make_cols:
        make_cols.append(col)
        next_col = col + 1
        while next_col < len(normalized_headers) and not normalized_headers[next_col]:
            make_cols.append(next_col)
            next_col += 1
    return sorted(set(make_cols))


def _first_row_value(row: tuple, columns: list[int]) -> str:
    for column in columns:
        if column < len(row) and row[column] not in (None, ""):
            return str(row[column]).strip()
    return ""


def _parse_make_list_rows(records: list[dict]) -> list[dict]:
    entries: list[dict] = []
    seen: set[tuple[str, str]] = set()
    for record in records:
        row = {key: _json_safe(value) for key, value in record["display_values"].items()}
        category = _pick(row, CATEGORY_KEYS)
        category = str(category).strip() if category else ""
        for make in _make_values_from_row(row):
            key = _make_entry_key(make, category)
            if key in seen:
                continue
            seen.add(key)
            entries.append({"make": make, "category": category})
    return entries


def _parse_make_list_pdf(file_path: str) -> list[dict]:
    """Extract make list from PDF using AI."""
    import pypdf
    from ai.service import AIService

    text = ""
    try:
        with open(file_path, "rb") as f:
            reader = pypdf.PdfReader(f)
            for page in reader.pages:
                text += (page.extract_text() or "") + "\n"
    except Exception:
        logger.exception("Failed to read PDF for make list.")
        return []

    if not text.strip():
        return []

    ai = AIService()
    if not ai.is_enabled():
        return []

    try:
        results = ai.run_json_prompt("extract_make_list.txt", text=text[:30000])
        entries = results.get("makes", results) if isinstance(results, dict) else results
        if not isinstance(entries, list):
            return []
        parsed: list[dict] = []
        seen: set[str] = set()
        for entry in entries:
            if not isinstance(entry, dict):
                continue
            make = str(entry.get("make") or "").strip()
            if not make or make.lower() in seen:
                continue
            seen.add(make.lower())
            parsed.append(
                {"make": make, "category": str(entry.get("category") or "").strip()}
            )
        return parsed
    except Exception:
        logger.exception("Failed to extract make list using AI.")
        return []
