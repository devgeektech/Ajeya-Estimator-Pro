import io
from django.test import TestCase
from unittest.mock import patch
from openpyxl import load_workbook
from apps.boq.models import BOQ
from apps.boq.services.boq_export_service import BOQExportService
from common.choices import BOQStatus
from django.contrib.auth import get_user_model

User = get_user_model()

class BOQExportTestCase(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(email="test@example.com", password="password")  # type: ignore
        boq_data = {
            "schema_version": 1,
            "headers": [
                {"key": "serial", "label": "Sr. No.", "index": 0},
                {"key": "description", "label": "Item Description", "index": 1},
                {"key": "qty", "label": "Qty", "index": 2},
                {"key": "unit", "label": "Unit", "index": 3},
                {"key": "rate", "label": "Rate", "index": 4},
                {"key": "amount", "label": "Amount", "index": 5},
            ],
            "rows": [
                {
                    "row_id": "r1",
                    "serial": "1",
                    "description": "Fire Pump Set",
                    "qty": 1,
                    "unit": "Set",
                    "sheet_name": "BOQ",
                    "excel_row_number": 2,
                }
            ]
        }
        
        # We need an analysis_data dict that pretends the pipeline is complete
        analysis_data = {
            "pricing_ready": True,
            "rows": [
                {
                    "row_id": "r1",
                    "products": [
                        {
                            "description_hint": "Fire Pump Set",
                            "category": "PUMP",
                            "status": "matched",
                            "review_output": {
                                "base_purchase_rate": 5000,
                                "qty": 1,
                                "final_material_amount": 5000,
                            }
                        }
                    ]
                }
            ]
        }

        self.boq = BOQ.objects.create(
            boq_name="Test BOQ Export",
            status=BOQStatus.MATCHING, # Doesn't strictly matter for the service class if we override validation
            boq_data=boq_data,
            analysis_data=analysis_data,
            user=self.user
        )

    @patch("apps.boq.services.boq_export_service.BOQReviewDisplayService")
    def test_boq_export_creates_workbook(self, mock_display):
        # Mock display data
        mock_display.return_value.build.return_value = {
            "has_analysis": True,
            "lines": [
                {
                    "row_id": "r1",
                    "serial": "1",
                    "description": "Fire Pump Set",
                    "qty": 1,
                    "unit": "Set",
                    "is_activity_only": False,
                    "products": [],
                    "rate": 5000,
                    "amount": 5000,
                }
            ]
        }
        
        # Override the BOQ status to allow export
        self.boq.status = BOQStatus.READY_EXPORT
        self.boq.save()

        service = BOQExportService(self.boq.id)
        
        # Generate the combined workbook
        payload, filename = service.run()
        
        # Verify the payload is a valid Excel file
        buffer = io.BytesIO(payload)
        wb = load_workbook(buffer)
        
        self.assertIn("Review", wb.sheetnames)
        self.assertIn("BOQ", wb.sheetnames)
        
        # Verify Review sheet headers
        review_ws = wb["Review"]
        self.assertEqual(review_ws["A1"].value, "Ser no of BOQ")
        self.assertEqual(review_ws["B1"].value, "BOQ Description")
        
        # Verify status transitioned
        self.boq.refresh_from_db()
        self.assertEqual(self.boq.status, BOQStatus.EXPORTED)


class BOQExportNotAvailableBorderTests(TestCase):
    def test_not_available_rate_amount_cells_have_borders(self):
        from openpyxl import Workbook

        from apps.boq.services.boq_export_service import (
            NOT_AVAILABLE_LABEL,
            _highlight_row,
            _write_not_available_rate_amount,
        )

        workbook = Workbook()
        sheet = workbook.active
        _write_not_available_rate_amount(sheet, 2, rate_col=5, amount_col=6)
        _highlight_row(sheet, 2, max_column=6)

        rate_cell = sheet.cell(row=2, column=5)
        amount_cell = sheet.cell(row=2, column=6)
        self.assertEqual(rate_cell.value, NOT_AVAILABLE_LABEL)
        self.assertEqual(amount_cell.value, NOT_AVAILABLE_LABEL)
        self.assertEqual(rate_cell.border.left.style, "thin")
        self.assertEqual(rate_cell.border.right.style, "thin")
        self.assertEqual(amount_cell.border.left.style, "thin")
        self.assertEqual(amount_cell.border.right.style, "thin")
        self.assertIsNotNone(rate_cell.fill)
        self.assertIsNotNone(amount_cell.fill)
